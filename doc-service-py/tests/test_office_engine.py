"""Real OfficeEngine tests: parse tiny .docx / .xlsx into canonical blocks.

Builds minimal Word and Excel files in memory (python-docx / openpyxl), runs
them through :class:`OfficeEngine` (and the :class:`ParserRouter` for the
integration path), and asserts the canonical-block invariants: block types,
heading breadcrumb, Vietnamese round-trip, the office field-value contract
(bbox / confidence / latex / page / parse_engine / parse_version), and the
graceful stub fallback when a backing library is unavailable.
"""

from __future__ import annotations

import io

import pytest

from app.engines import office_parser
from app.engines.base import PageKind, ParseContext
from app.engines.stub_engines import OfficeEngine
from app.normalizer import normalize
from app.router import ParserRouter

_DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _docx_bytes() -> bytes:
    """Build a tiny .docx with headings, a paragraph, an empty para, a table."""
    import docx  # noqa: PLC0415

    document = docx.Document()
    document.add_heading("Tài liệu thử nghiệm", level=1)
    document.add_heading("Mục con", level=2)
    document.add_paragraph("Nội dung tiếng Việt có dấu.")
    document.add_paragraph("")  # empty paragraph -> must be skipped
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Khoá"
    table.cell(0, 1).text = "Giá trị"
    table.cell(1, 0).text = "Alpha"
    table.cell(1, 1).text = "1"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    """Build a tiny .xlsx with two sheets, Vietnamese header + data."""
    import openpyxl  # noqa: PLC0415

    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Bảng 1"
    ws1.append(["Tên", "Điểm"])
    ws1.append(["Học viên A", 9])
    ws1.append(["Học viên B", 8])
    ws2 = wb.create_sheet("Bảng 2")
    ws2.append(["Ghi chú"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _ctx(mime: str, filename: str) -> ParseContext:
    return ParseContext(
        job_id="job-office",
        document_id="doc-office-1",
        correlation_id="corr-1",
        filename=filename,
        mime=mime,
        page=1,
        order_start=0,
    )


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def test_office_parses_docx_into_canonical_blocks():
    pytest.importorskip("docx")
    data = _docx_bytes()
    blocks = OfficeEngine().parse(data, _ctx(_DOCX_MIME, "test.docx"))

    assert len(blocks) >= 4  # 2 headings + 1 paragraph + 1 table (empty skipped)

    # First block: Heading 1 with the breadcrumb seeded.
    first = blocks[0]
    assert first.type == "heading"
    assert first.heading_path == ["Tài liệu thử nghiệm"]

    # Heading 2 appends under Heading 1.
    second = blocks[1]
    assert second.type == "heading"
    assert second.heading_path == ["Tài liệu thử nghiệm", "Mục con"]

    # A paragraph exists with the Vietnamese text round-tripping exactly.
    paragraphs = [b for b in blocks if b.type == "paragraph"]
    assert any(b.text == "Nội dung tiếng Việt có dấu." for b in paragraphs)
    assert paragraphs[0].heading_path == ["Tài liệu thử nghiệm", "Mục con"]

    # A table block exists with markdown + html rendering.
    tables = [b for b in blocks if b.type == "table"]
    assert tables, "expected a table block"
    assert "|" in tables[0].markdown
    assert "<table" in tables[0].html
    assert "Khoá" in tables[0].text and "Giá trị" in tables[0].text

    # Empty paragraph produced no block (count stays at the 4 real items).
    assert len(blocks) == 4

    # Office field-value contract on every block.
    for b in blocks:
        assert b.parse_engine == "office"
        assert b.bbox == (0.0, 0.0, 0.0, 0.0)
        assert b.confidence == 1.0
        assert b.latex is None
        assert b.page == 1
        assert b.parse_version.startswith("python-docx")


def test_office_docx_routes_through_parser_router():
    pytest.importorskip("docx")
    data = _docx_bytes()
    router = ParserRouter()

    decisions = router.classify(_DOCX_MIME, data)
    assert len(decisions) == 1
    assert decisions[0].kind == PageKind.OFFICE

    engine = router.engine_for(PageKind.OFFICE)
    assert isinstance(engine, OfficeEngine)
    assert engine.supports(PageKind.OFFICE) is True

    blocks = engine.parse(data, _ctx(_DOCX_MIME, "routed.docx"))
    assert blocks and blocks[0].type == "heading"
    assert all(b.parse_engine == "office" for b in blocks)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def test_office_parses_xlsx_into_canonical_blocks():
    pytest.importorskip("openpyxl")
    data = _xlsx_bytes()
    blocks = OfficeEngine().parse(data, _ctx(_XLSX_MIME, "test.xlsx"))

    # First block per sheet is a heading with the sheet name.
    assert blocks[0].type == "heading"
    assert blocks[0].text == "Bảng 1"
    assert blocks[0].heading_path == ["Bảng 1"]

    # The block right after the sheet heading is the sheet's table.
    assert blocks[1].type == "table"
    assert blocks[1].heading_path == ["Bảng 1"]

    # Second sheet "Bảng 2" appears as its own heading block.
    headings = [b for b in blocks if b.type == "heading"]
    assert any(b.text == "Bảng 2" for b in headings)

    # Vietnamese cell text round-trips inside the table block.
    tables = [b for b in blocks if b.type == "table"]
    assert any("Học viên A" in b.text for b in tables)
    assert any("Điểm" in b.text for b in tables)

    for b in blocks:
        assert b.parse_engine == "office"
        assert b.parse_version.startswith("openpyxl")
        assert b.bbox == (0.0, 0.0, 0.0, 0.0)
        assert b.confidence == 1.0
        assert b.latex is None
        assert b.page == 1


# --------------------------------------------------------------------------- #
# Fallback + normalization
# --------------------------------------------------------------------------- #
def test_office_falls_back_to_stub_when_docx_lib_unavailable(monkeypatch):
    # Force the lazy loader to report python-docx as absent.
    monkeypatch.setattr(office_parser, "load_docx", lambda: None)

    # Any bytes are fine; the fallback happens before parsing.
    blocks = OfficeEngine().parse(b"not-a-real-docx", _ctx(_DOCX_MIME, "x.docx"))

    # Stub fallback shape: exactly 2 blocks, first a heading. Must not raise.
    assert len(blocks) == 2
    assert blocks[0].type == "heading"
    assert all(b.parse_engine == "office" for b in blocks)


def test_office_blocks_normalize_into_valid_document():
    pytest.importorskip("docx")
    data = _docx_bytes()
    blocks = OfficeEngine().parse(data, _ctx(_DOCX_MIME, "norm.docx"))

    document = normalize(
        document_id="doc-office-1",
        filename="norm.docx",
        sha256="0" * 64,
        mime=_DOCX_MIME,
        pages=1,
        page_blocks=[blocks],
    )

    # Global reading order is re-assigned monotonically 0..n-1.
    orders = [b.order for b in document.blocks]
    assert orders == list(range(len(document.blocks)))
    assert document.stats.engine_breakdown.get("office") == len(document.blocks)
